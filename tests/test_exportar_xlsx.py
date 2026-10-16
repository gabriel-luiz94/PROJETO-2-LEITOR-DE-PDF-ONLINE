"""TASK-070 — exportação .xlsx (planilhas ODI e ODD) do Resultado do Orçamento."""
import io

import openpyxl
import pytest
from fastapi.testclient import TestClient

import database
from middleware.auth_middleware import create_jwt_token
from services.exportar_xlsx import ErroExportacao, gerar_xlsx

ODI = [{"mdo": "MATERIAL", "codigo": "000123", "desc": "CABO CAA 2", "qtd": 12.5},
       {"mdo": "MAO-DE-OBRA", "codigo": "9", "desc": "INSTALAR", "qtd": "1,00"}]
ODD = [{"mdo": "M", "codigo": "55", "desc": "POSTE", "qtd": 3}]


def abrir(conteudo):
    return openpyxl.load_workbook(io.BytesIO(conteudo))


def test_duas_planilhas_odi_e_odd_com_cabecalho():
    wb = abrir(gerar_xlsx(ODI, ODD))
    assert wb.sheetnames == ["ODI", "ODD"]
    assert [c.value for c in wb["ODI"][1]] == ["TIPO MDO", "CÓDIGO", "QUANTIDADE", "CENTRO", "DESCRIÇÃO"]
    assert [[c.value for c in r] for r in wb["ODI"].iter_rows(min_row=2)] == [["M", "000123", 12.5, "1600", "CABO CAA 2"], ["N", "9", 1.0, "1000", "INSTALAR"]]
    assert [[c.value for c in r] for r in wb["ODD"].iter_rows(min_row=2)] == [["M", "55", 3.0, "1600", "POSTE"]]


def test_codigo_com_zeros_a_esquerda_continua_texto_e_quantidade_numerica():
    ws = abrir(gerar_xlsx(ODI, [])) ["ODI"]
    assert ws["B2"].value == "000123" and ws["B2"].data_type == "s"
    assert isinstance(ws["C2"].value, float)


def test_texto_que_parece_formula_nao_vira_formula():
    ws = abrir(gerar_xlsx([{"mdo": "M", "codigo": "1", "desc": "=1+1", "qtd": 1}], []))["ODI"]
    assert ws["E2"].value == "=1+1" and ws["E2"].data_type == "s"


def test_planilhas_vazias_so_com_cabecalho_e_caracteres_ilegais_removidos():
    wb = abrir(gerar_xlsx([{"mdo": "M", "codigo": "1\x00", "desc": "a\x01b", "qtd": 1}], []))
    assert wb["ODD"].max_row == 1 and wb["ODI"]["E2"].value == "ab"


def test_limites_e_formato_invalido():
    with pytest.raises(ErroExportacao):
        gerar_xlsx("não é lista", [])
    with pytest.raises(ErroExportacao):
        gerar_xlsx([{}] * 20001, [])
    with pytest.raises(ErroExportacao):
        gerar_xlsx([1], [])


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(database, "DB_PATH", str(tmp_path / "teste.db"))
    database.init_db()
    import app as appmod
    return TestClient(appmod.app)


def test_rota_devolve_xlsx(client):
    h = {"Authorization": f"Bearer {create_jwt_token('u', 'u@x.com', 'operador')}"}
    r = client.post("/api/orcamento/exportar-xlsx", json={"odi": ODI, "odd": ODD}, headers=h)
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert abrir(r.content).sheetnames == ["ODI", "ODD"]
    assert client.post("/api/orcamento/exportar-xlsx", json={"odi": "x"}, headers=h).status_code == 400
    assert client.post("/api/orcamento/exportar-xlsx", content=b"nao-json", headers=h).status_code == 400
    assert client.post("/api/orcamento/exportar-xlsx", json={"odi": ODI}).status_code == 401
