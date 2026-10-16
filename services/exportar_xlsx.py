"""
services/exportar_xlsx.py — Exportação do Resultado do Orçamento no formato SAP (TASK-070): um .xlsx com duas planilhas,
ODI (tudo que está instalando) e ODD (tudo que está removendo), colunas TIPO MDO · CÓDIGO · QUANTIDADE · CENTRO (1600 material / 1000 mão de obra) · DESCRIÇÃO.
"""
import io
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

LIMITE_LINHAS = 20000
LIMITE_TEXTO = 500
CABECALHO = ["TIPO MDO", "CÓDIGO", "QUANTIDADE", "CENTRO", "DESCRIÇÃO"]
CENTRO = {"M": "1600", "N": "1000"}                        # material = 1600; mão de obra = 1000 (derivado do TIPO MDO)
_ILEGAIS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")        # caracteres que o XML do Excel não aceita
_MDO = {"M": "M", "MATERIAL": "M", "N": "N", "S": "N", "MAO-DE-OBRA": "N", "MÃO-DE-OBRA": "N", "": "N", "-": "N"}


class ErroExportacao(ValueError):
    pass


def _texto(v) -> str:
    return _ILEGAIS.sub("", "" if v is None else str(v)).strip()[:LIMITE_TEXTO]


def _qtd(v) -> float:
    if isinstance(v, bool):
        return 0.0
    if isinstance(v, (int, float)):
        return round(float(v), 2)
    try:
        return round(float(str(v).strip().replace(".", "").replace(",", ".")), 2) if str(v).strip() else 0.0
    except ValueError:
        return 0.0


def _validar(nome: str, linhas) -> list:
    if not isinstance(linhas, list):
        raise ErroExportacao(f"{nome}: precisa ser uma lista.")
    if len(linhas) > LIMITE_LINHAS:
        raise ErroExportacao(f"{nome}: linhas demais (máx. {LIMITE_LINHAS}).")
    for i, x in enumerate(linhas, 1):
        if not isinstance(x, dict):
            raise ErroExportacao(f"{nome} linha {i}: precisa ser um objeto.")
    return linhas


def _planilha(ws, linhas):
    ws.append(CABECALHO)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(horizontal="center")
    for x in linhas:
        mdo = _MDO.get(_texto(x.get("mdo")).upper(), _texto(x.get("mdo")).upper()[:1] or "N")
        ws.append([mdo, _texto(x.get("codigo")), _qtd(x.get("qtd")), CENTRO.get(mdo, ""), _texto(x.get("desc"))])
        for c in ws[ws.max_row]:
            if isinstance(c.value, str) and c.value[:1] in ("=", "+", "-", "@"):
                c.data_type = "s"                                  # texto, nunca fórmula
        ws.cell(row=ws.max_row, column=3).number_format = "#,##0.00"
    for col, larg in zip("ABCDE", (10, 16, 14, 10, 60)):
        ws.column_dimensions[col].width = larg
    ws.freeze_panes = "A2"


def gerar_xlsx(odi, odd) -> bytes:
    """Bytes de um .xlsx com as planilhas ODI e ODD. Cada linha: {mdo, codigo, desc, qtd}."""
    odi, odd = _validar("ODI", odi), _validar("ODD", odd)
    wb = Workbook()
    ws = wb.active
    ws.title = "ODI"
    _planilha(ws, odi)
    _planilha(wb.create_sheet("ODD"), odd)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
